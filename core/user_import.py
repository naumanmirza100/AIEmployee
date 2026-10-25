"""Adding a company's people from an Excel sheet.

Asked for after the client test of 25 Oct 2026: one row per person (name, email, phone, role and the HR details
such as the joining date), a password made for each, and the passwords shown once so they can be handed out.
Each row goes through the same checks as the Add User form (`api.views.company_users._new_login`); a refused row
is reported with its reason and the others are still added.
"""
import csv
import io
import re
import secrets
from datetime import date, datetime

from django.db import transaction

#: The sheet's columns, in order: (key, heading, required, what to write).
COLUMNS = [
    ('full_name', 'Full name', True, 'Letters only, e.g. Ayesha Khan'),
    ('email', 'Email', True, 'Their sign-in email; must not be used by anyone yet'),
    ('phone_number', 'Phone number', True, 'At least 7 digits, e.g. +92 300 1234567'),
    ('role', 'Role', True, 'Project Manager, Team Member, Developer, Viewer, Internee or Designer'),
    ('joining_date', 'Joining date', False, 'YYYY-MM-DD, e.g. 2026-11-01'),
    ('job_title', 'Job title', False, 'e.g. Flutter Developer'),
    ('department', 'Department', False, 'Made if the company does not have it yet'),
    ('employment_type', 'Employment type', False, 'Full-time, Part-time, Contractor or Intern'),
    ('date_of_birth', 'Date of birth', False, 'YYYY-MM-DD'),
    ('personal_email', 'Personal email', False, ''),
    ('manager_email', 'Manager email', False, 'The manager must already be in the company'),
    ('location', 'Location', False, 'e.g. Lahore'),
    ('bio', 'Bio', False, 'At least 10 letters or digits'),
    ('username', 'Username', False, 'Left empty: the part of the email before the @'),
]

ROLES = {'project manager': 'project_manager', 'team member': 'team_member', 'developer': 'developer',
         'viewer': 'viewer', 'internee': 'internee', 'intern': 'internee', 'designer': 'designer'}
EMPLOYMENT_TYPES = {'full-time': 'full_time', 'full time': 'full_time', 'part-time': 'part_time',
                    'part time': 'part_time', 'contractor': 'contract', 'contract': 'contract', 'intern': 'intern'}

# Headings people write instead of ours.
ALIASES = {'name': 'full_name', 'full name': 'full_name', 'phone': 'phone_number', 'mobile': 'phone_number',
           'phone no': 'phone_number', 'start date': 'joining_date', 'join date': 'joining_date',
           'joined': 'joining_date', 'title': 'job_title', 'designation': 'job_title', 'dob': 'date_of_birth',
           'manager': 'manager_email', 'type': 'employment_type'}

MAX_ROWS = 500


class SheetError(ValueError):
    """The file as a whole cannot be read; said in words for the screen."""


# ---------------------------------------------------------------------------
# Passwords
# ---------------------------------------------------------------------------

# Letters and digits that cannot be mistaken for each other when read out or copied from paper.
_LOWER = 'abcdefghjkmnpqrstuvwxyz'
_UPPER = 'ABCDEFGHJKLMNPQRSTUVWXYZ'
_DIGITS = '23456789'
_SPECIAL = '!@#$%&*?'


def make_password(length=12):
    """A password the Add User form accepts: upper and lower case, a digit and a special character."""
    picks = [secrets.choice(_LOWER), secrets.choice(_UPPER), secrets.choice(_DIGITS), secrets.choice(_SPECIAL)]
    every = _LOWER + _UPPER + _DIGITS + _SPECIAL
    picks += [secrets.choice(every) for _ in range(length - len(picks))]
    secrets.SystemRandom().shuffle(picks)
    return ''.join(picks)


# ---------------------------------------------------------------------------
# Reading the sheet
# ---------------------------------------------------------------------------

def _heading_key(text):
    clean = re.sub(r'\(.*?\)|\*', '', str(text or '')).strip().lower()
    clean = re.sub(r'\s+', ' ', clean)
    for key, heading, *_ in COLUMNS:
        if clean in (heading.lower(), key.replace('_', ' ')):
            return key
    return ALIASES.get(clean)


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)                          # a phone number typed into a number cell
    if isinstance(value, (datetime, date)):
        return (value.date() if isinstance(value, datetime) else value).isoformat()
    return str(value).strip()


def _rows_of(upload):
    """[[cell, ...], ...] from an .xlsx or .csv upload."""
    name = (getattr(upload, 'name', '') or '').lower()
    data = upload.read()
    if name.endswith('.csv'):
        text = data.decode('utf-8-sig', errors='replace')
        return list(csv.reader(io.StringIO(text)))
    if not name.endswith(('.xlsx', '.xlsm')):
        raise SheetError('Upload the template as an Excel file (.xlsx) or a CSV file.')
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise SheetError('This file could not be opened as an Excel sheet.') from exc
    sheet = book['Users'] if 'Users' in book.sheetnames else book.worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def read_sheet(upload):
    """[(row_number, {key: text})] for every filled row, with the sheet's row numbers."""
    rows = _rows_of(upload)
    header_at = next((i for i, row in enumerate(rows[:10]) if any(_heading_key(c) for c in row)), None)
    if header_at is None:
        raise SheetError('No column headings found. Use the template: Full name, Email, Phone number, Role …')
    keys = [_heading_key(c) for c in rows[header_at]]
    missing = [heading for key, heading, required, _ in COLUMNS if required and key not in keys]
    if missing:
        raise SheetError('The sheet has no column for: ' + ', '.join(missing) + '.')
    people = []
    for offset, row in enumerate(rows[header_at + 1:], start=header_at + 2):
        values = {key: _cell_text(cell) for key, cell in zip(keys, row) if key}
        if any(values.values()):
            people.append((offset, values))
    if len(people) > MAX_ROWS:
        raise SheetError(f'The sheet has {len(people)} people; add at most {MAX_ROWS} at a time.')
    return people


def _date(text):
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    raise ValueError(f'"{text}" is not a date; write it as YYYY-MM-DD.')


# ---------------------------------------------------------------------------
# Adding the people
# ---------------------------------------------------------------------------

def _hr_details(company, user, values):
    """Put the sheet's HR columns on the person's HR record (made with their login). Returns notes for the row."""
    from hr_agent.models import Department, Employee

    record = Employee.objects.filter(company=company, user=user).first()
    if record is None:
        return ['No HR record was made for this login, so its HR details were not saved.']
    notes, fields = [], []
    for key, attr in (('joining_date', 'start_date'), ('date_of_birth', 'date_of_birth')):
        if values.get(key):
            setattr(record, attr, _date(values[key]))
            fields.append(attr)
    for key, attr, limit in (('job_title', 'job_title', 160), ('personal_email', 'personal_email', 254)):
        if values.get(key):
            setattr(record, attr, values[key][:limit])
            fields.append(attr)
    if values.get('employment_type'):
        kind = EMPLOYMENT_TYPES.get(values['employment_type'].strip().lower())
        if kind:
            record.employment_type = kind
            fields.append('employment_type')
        else:
            notes.append(f'Employment type "{values["employment_type"]}" is not one of Full-time, Part-time, '
                         'Contractor, Intern; left as Full-time.')
    if values.get('department'):
        name = values['department'][:120]
        record.department_obj = (Department.objects.filter(company=company, name__iexact=name).first()
                                 or Department.objects.create(company=company, name=name))
        record.department = record.department_obj.name
        fields += ['department_obj', 'department']
    if values.get('manager_email'):
        manager = (Employee.objects.filter(company=company, work_email__iexact=values['manager_email'])
                   .exclude(pk=record.pk).first())
        if manager:
            record.manager = manager
            fields.append('manager')
        else:
            notes.append(f'No one in the company has the email {values["manager_email"]}; no manager set.')
    if fields:
        record.save(update_fields=sorted(set(fields)) + ['updated_at'])
    return notes


def add_people(company_user, people):
    """Add each person as the Add User form would, with a password made for them.

    Returns one result per row: {row, full_name, email, role, status: 'added'|'refused', password, reason,
    notes}. The passwords are only ever in this result."""
    from rest_framework.response import Response
    from api.views.company_users import _new_login

    labels = {value: label.title() for label, value in ROLES.items() if label != 'intern'}
    results, seen = [], set()
    for row_number, values in people:
        email = values.get('email', '').strip().lower()
        result = {'row': row_number, 'full_name': values.get('full_name', ''), 'email': email,
                  'role': values.get('role', ''), 'status': 'refused', 'password': '', 'reason': '', 'notes': []}
        results.append(result)
        if email and email in seen:
            result['reason'] = 'This email is on the sheet twice; only the first row was added.'
            continue
        seen.add(email)
        role = ROLES.get(values.get('role', '').strip().lower())
        if role is None:
            result['reason'] = ('Role must be one of Project Manager, Team Member, Developer, Viewer, '
                                'Internee or Designer.')
            continue
        try:
            for key in ('joining_date', 'date_of_birth'):
                _date(values.get(key))
        except ValueError as exc:
            result['reason'] = str(exc)
            continue
        password = make_password()
        form = {'fullName': values.get('full_name', ''), 'email': email, 'password': password,
                'phoneNumber': values.get('phone_number', ''), 'role': role,
                'location': values.get('location', ''), 'bio': values.get('bio', ''),
                'username': values.get('username', '')}
        try:
            with transaction.atomic():
                made = _new_login(company_user, form)
                if isinstance(made, Response):
                    result['reason'] = (made.data or {}).get('message') or 'Refused.'
                    transaction.set_rollback(True)
                    continue
                user = made[0]
                result['notes'] = _hr_details(company_user.company, user, values)
        except Exception as exc:          # one bad row must not stop the rest
            result['reason'] = f'Could not be added: {exc}'[:300]
            continue
        result.update(status='added', password=password, role=labels.get(role, role), username=user.username)
    return results


# ---------------------------------------------------------------------------
# The template
# ---------------------------------------------------------------------------

def template_workbook():
    """The .xlsx to fill in: a Users sheet with the headings, dropdowns for Role and Employment type, and a sheet
    saying what goes in each column."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    users = book.active
    users.title = 'Users'
    users.append([heading + (' *' if required else '') for _, heading, required, _ in COLUMNS])
    bold, fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='5B21B6')
    for index, (_key, heading, _required, _note) in enumerate(COLUMNS, start=1):
        cell = users.cell(row=1, column=index)
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(vertical='center')
        users.column_dimensions[cell.column_letter].width = max(16, len(heading) + 6)
    users.freeze_panes = 'A2'

    letters = {key: users.cell(row=1, column=i).column_letter for i, (key, *_r) in enumerate(COLUMNS, start=1)}
    role_list = DataValidation(type='list', allow_blank=False, showDropDown=False,
                               formula1='"Project Manager,Team Member,Developer,Viewer,Internee,Designer"')
    type_list = DataValidation(type='list', allow_blank=True, showDropDown=False,
                               formula1='"Full-time,Part-time,Contractor,Intern"')
    users.add_data_validation(role_list)
    users.add_data_validation(type_list)
    role_list.add(f'{letters["role"]}2:{letters["role"]}{MAX_ROWS + 1}')
    type_list.add(f'{letters["employment_type"]}2:{letters["employment_type"]}{MAX_ROWS + 1}')
    for key in ('joining_date', 'date_of_birth'):
        for row in range(2, MAX_ROWS + 2):
            users[f'{letters[key]}{row}'].number_format = 'yyyy-mm-dd'
    for key in ('phone_number',):
        for row in range(2, MAX_ROWS + 2):
            users[f'{letters[key]}{row}'].number_format = '@'          # text, so a leading + or 0 stays

    guide = book.create_sheet('How to fill')
    guide.append(['Column', 'Required', 'What to write'])
    for cell in guide[1]:
        cell.font, cell.fill = bold, fill
    for _key, heading, required, note in COLUMNS:
        guide.append([heading, 'Yes' if required else '', note])
    guide.append([])
    guide.append(['Passwords', '', 'Not in the sheet: one is made for each person and shown once after the '
                                   'upload, to hand out. Each person can change theirs after signing in.'])
    guide.column_dimensions['A'].width = 20
    guide.column_dimensions['B'].width = 10
    guide.column_dimensions['C'].width = 90

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
