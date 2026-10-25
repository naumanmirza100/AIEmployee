"""A company adds its people from an Excel sheet, each with a password made for them.

Asked for after the client test of 25 Oct 2026: name, email, phone, role and the HR details on one row per
person, the passwords shown once to hand out, and a template sheet to start from.
"""
import io
from datetime import date

from django.contrib.auth import authenticate
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from api.views import company_users
from core.models import Company, CompanyUser
from core.user_import import COLUMNS, make_password
from hr_agent.models import Employee

HEADINGS = ['Full name *', 'Email *', 'Phone number *', 'Role *', 'Joining date', 'Job title', 'Department',
            'Employment type', 'Manager email']


def sheet(*rows, headings=HEADINGS):
    book = Workbook()
    users = book.active
    users.title = 'Users'
    users.append(headings)
    for row in rows:
        users.append(row)
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile('people.xlsx', out.getvalue(),
                             content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class UserImportTests(TestCase):

    def setUp(self):
        self.company = Company.objects.create(name='Acme', email='acme@test.local')
        self.admin = CompanyUser.objects.create(company=self.company, email='admin@acme.test', full_name='Ann Admin',
                                                role='admin', password_hash='x', is_active=True)
        Employee.objects.create(company=self.company, full_name='Bilal Boss', work_email='boss@acme.test')

    def upload(self, file):
        request = APIRequestFactory().post('/x', {'file': file}, format='multipart')
        force_authenticate(request, user=self.admin)
        response = company_users.users_import(request)
        return response.status_code, response.data

    def test_each_person_is_added_with_a_password_and_their_hr_details(self):
        code, body = self.upload(sheet(
            ['Sara Malik', 'sara@acme.test', '+92 300 1234567', 'Developer', date(2026, 11, 2), 'Flutter Developer',
             'Mobile', 'Full-time', 'boss@acme.test'],
            ['Omar Raza', 'omar@acme.test', '03001234568', 'Project Manager', '2026-11-09', '', 'Delivery',
             'Contractor', '']))
        self.assertEqual(code, 200, body)
        self.assertEqual((body['data']['added'], body['data']['refused']), (2, 0))
        sara = body['data']['rows'][0]
        self.assertEqual((sara['status'], sara['role'], sara['row']), ('added', 'Developer', 2))
        user = User.objects.get(email='sara@acme.test')
        self.assertEqual(authenticate(username=user.username, password=sara['password']), user)
        record = Employee.objects.get(user=user)
        self.assertEqual((record.start_date, record.job_title, record.department, record.employment_type,
                          record.manager.full_name),
                         (date(2026, 11, 2), 'Flutter Developer', 'Mobile', 'full_time', 'Bilal Boss'))
        self.assertEqual(Employee.objects.get(user__email='omar@acme.test').employment_type, 'contract')

    def test_a_bad_row_is_refused_with_its_reason_and_the_rest_are_added(self):
        code, body = self.upload(sheet(
            ['Sara Malik', 'sara@acme.test', '+92 300 1234567', 'Developer'],
            ['Sara Again', 'sara@acme.test', '+92 300 1234569', 'Developer'],
            ['Nida 2', 'nida@acme.test', '+92 300 1234570', 'Developer'],
            ['Zain Ali', 'zain@acme.test', '+92 300 1234571', 'Boss'],
            ['Hira Khan', 'hira@acme.test', '12', 'Viewer'],
            ['Asad Noor', 'asad@acme.test', '+92 300 1234572', 'Viewer', 'next week']))
        rows = {row['email']: row for row in body['data']['rows'] if row['email'] != 'sara@acme.test'}
        self.assertEqual(body['data']['added'], 1)
        self.assertIn('twice', body['data']['rows'][1]['reason'])
        self.assertIn('must not contain numbers', rows['nida@acme.test']['reason'])
        self.assertIn('Role must be one of', rows['zain@acme.test']['reason'])
        self.assertIn('valid phone number', rows['hira@acme.test']['reason'])
        self.assertIn('is not a date', rows['asad@acme.test']['reason'])
        self.assertEqual(rows['zain@acme.test']['password'], '')
        self.assertFalse(User.objects.filter(email__in=['nida@acme.test', 'zain@acme.test', 'hira@acme.test']).exists())

    def test_a_csv_works_too(self):
        text = 'Full name,Email,Phone number,Role\nSara Malik,sara@acme.test,+92 300 1234567,Team Member\n'
        code, body = self.upload(SimpleUploadedFile('people.csv', text.encode(), content_type='text/csv'))
        self.assertEqual(body['data']['added'], 1, body)

    def test_a_sheet_without_the_required_columns_is_refused_whole(self):
        code, body = self.upload(sheet(['Sara', 'sara@acme.test'], headings=['Full name', 'Email']))
        self.assertEqual(code, 400)
        self.assertIn('no column for: Phone number, Role', body['message'])

    def test_the_template_has_every_column_and_a_role_dropdown(self):
        request = APIRequestFactory().get('/x')
        force_authenticate(request, user=self.admin)
        response = company_users.users_import_template(request)
        book = load_workbook(io.BytesIO(response.content))
        headings = [cell.value for cell in book['Users'][1]]
        self.assertEqual(len(headings), len(COLUMNS))
        self.assertEqual(headings[:4], ['Full name *', 'Email *', 'Phone number *', 'Role *'])
        rules = [dv.formula1 for dv in book['Users'].data_validations.dataValidation]
        self.assertIn('"Project Manager,Team Member,Developer,Viewer,Internee,Designer"', rules)
        self.assertIn('How to fill', book.sheetnames)

    def test_made_passwords_pass_the_add_user_rules(self):
        import re
        for _ in range(50):
            password = make_password()
            self.assertEqual(len(password), 12)
            for rule in (r'[A-Z]', r'[a-z]', r'[0-9]', r'[!@#$%&*?]'):
                self.assertRegex(password, rule)
