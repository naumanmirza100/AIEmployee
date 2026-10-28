"""The Operations gaps of the 10 Oct sweep (MDS/01-Plan/gaps/AGENTS_INSIDE_GAPS_REPORT.md, CO-172 to CO-279), fixed
on 28 Oct 2026 before the Operations demo. One test class per item; the AI is stubbed everywhere."""
import json
import os
import tempfile
from datetime import timedelta
from unittest import mock

from django.test import Client, TestCase
from django.utils import timezone

from core.models import Company, CompanyModulePurchase, CompanyUser, CompanyUserToken
from operations_agent import intent, tasks
from operations_agent.agents import document_processing_agent as dpa
from operations_agent.agents.document_processing_agent import DocumentProcessingAgent
from operations_agent.agents.summarization_agent import DocumentSummarizationAgent, clean_deadlines
from operations_agent.chunking import looks_like_toc_or_index
from operations_agent.models import (
    OperationsDocument, OperationsDocumentChunk, OperationsDocumentSummary, OperationsGeneratedDocument,
)


def a_file(suffix, data=b'', text=None):
    handle, path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(handle, 'wb') as f:
        f.write(text.encode('utf-8') if text is not None else data)
    return path


class OpsCase(TestCase):

    def setUp(self):
        self.company = Company.objects.create(name='Margalla', email='m@test.local')
        CompanyModulePurchase.objects.create(company=self.company, module_name='operations_agent', status='active',
                                             is_complimentary=True)
        self.nadia = CompanyUser.objects.create(company=self.company, email='nadia@test.local', full_name='Nadia',
                                                role='admin', password_hash='x', is_active=True)
        key = CompanyUserToken.objects.get_or_create(company_user=self.nadia)[0].key
        self.screen = Client(HTTP_AUTHORIZATION=f'Token {key}')

    def placeholder(self, **extra):
        return OperationsDocument.objects.create(**{'company': self.company, 'uploaded_by': self.nadia,
                                                    'title': 'Lease', 'original_filename': 'lease.txt', 'file': '',
                                                    'processing_status': 'pending', **extra})

    def quiet_agent(self):
        """A processing agent whose three AI steps answer without calling anything."""
        agent = DocumentProcessingAgent()
        agent._classify_document = mock.Mock(return_value={'document_type': 'contract'})
        agent._extract_entities = mock.Mock(return_value={'entities': {}})
        agent._generate_summary = mock.Mock(return_value={'summary': 'A lease.', 'key_insights': []})
        return agent


class DeletedWhileProcessingTests(OpsCase):
    """CO-172: a document deleted during its AI steps came back, searchable."""

    def test_it_does_not_come_back(self):
        doc = self.placeholder()
        agent = self.quiet_agent()

        def delete_it(**kwargs):
            OperationsDocument.objects.filter(pk=doc.pk).delete()      # the bin icon, during the AI steps
            return {'summary': 'A lease.', 'key_insights': []}
        agent._generate_summary = mock.Mock(side_effect=delete_it)
        path = a_file('.txt', text='The rent is PKR 40,000 a month, paid by the 5th. ' * 20)
        with mock.patch.object(dpa, 'build_embedding_service', return_value=None):
            result = agent._process_file(path, 'lease.txt', self.company.id, self.nadia.id, existing_doc_id=doc.pk)
        self.assertTrue(result.get('deleted'))
        self.assertFalse(OperationsDocument.objects.filter(pk=doc.pk).exists())
        self.assertFalse(OperationsDocumentChunk.objects.exists())
        self.assertFalse(os.path.exists(path))

    def test_deleted_before_it_started_makes_nothing(self):
        path = a_file('.txt', text='Hello')
        result = self.quiet_agent()._process_file(path, 'x.txt', self.company.id, self.nadia.id, existing_doc_id=987)
        self.assertTrue(result.get('deleted'))
        self.assertEqual(OperationsDocument.objects.count(), 0)


class FailureNotDiscardedTests(OpsCase):
    """CO-178: an error handed back as a result left the entry on 'processing' for half an hour."""

    def test_the_entry_is_failed_with_the_agents_error(self):
        doc = self.placeholder(processing_status='processing')
        tasks._fail_if_still_running(OperationsDocument, doc.pk, {'success': False, 'error': 'Disk full'})
        doc.refresh_from_db()
        self.assertEqual((doc.processing_status, doc.processing_error), ('failed', 'Disk full'))

    def test_a_failure_with_its_own_message_is_left_alone(self):
        doc = self.placeholder(processing_status='failed', processing_error='Scanned PDF')
        tasks._fail_if_still_running(OperationsDocument, doc.pk, {'success': False, 'error': 'other'})
        doc.refresh_from_db()
        self.assertEqual(doc.processing_error, 'Scanned PDF')


class WholeFileReadTests(TestCase):
    """CO-177: tables in Word and PowerPoint, grouped shapes and speaker notes were left out."""

    def test_a_word_table_is_read_in_its_place(self):
        from docx import Document
        word = Document()
        word.add_paragraph('Fee schedule below.')
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'Service', 'Fee'
        table.cell(1, 0).text, table.cell(1, 1).text = 'Hosting', 'PKR 18,500'
        word.add_paragraph('Payment within 30 days.')
        path = a_file('.docx')
        word.save(path)
        ok, text, _pages, _err = DocumentProcessingAgent()._extract_text(path, 'docx')
        self.assertTrue(ok)
        self.assertIn('Hosting | PKR 18,500', text)
        self.assertLess(text.index('Fee schedule'), text.index('Hosting'))
        self.assertLess(text.index('Hosting'), text.index('Payment within'))

    def test_a_slide_table_and_its_notes_are_read(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        shape.table.cell(0, 0).text, shape.table.cell(0, 1).text = 'Quarter', 'Revenue'
        shape.table.cell(1, 0).text, shape.table.cell(1, 1).text = 'Q3', 'PKR 2.4M'
        slide.notes_slide.notes_text_frame.text = 'Say the Q3 number twice.'
        path = a_file('.pptx')
        deck.save(path)
        agent = DocumentProcessingAgent()
        ok, text, _pages, _err = agent._extract_text(path, 'pptx')
        self.assertTrue(ok)
        self.assertIn('Q3 | PKR 2.4M', text)
        self.assertIn('Speaker notes: Say the Q3 number twice.', text)
        self.assertEqual(agent.page_starts[0][1], 1)                    # slide 1 starts the text (CO-278)


class AnyScriptTests(TestCase):
    """CO-215: Urdu-script passages and number-only stretches were thrown away as junk."""

    def test_urdu_and_figures_are_content(self):
        self.assertFalse(looks_like_toc_or_index('یہ معاہدہ دونوں فریقین کے درمیان ایک سال کے لیے ہے اور کرایہ ہر ماہ کی پانچ تاریخ تک ادا کیا جائے گا'))
        self.assertFalse(looks_like_toc_or_index('2026-01-01  18500  23000  41500\n2026-02-01  19100  22400  41500\n'
                                                 '2026-03-01  20500  21000  41500'))
        self.assertTrue(looks_like_toc_or_index('Contents ........ 3\nScope ........ 4\nTerms ........ 7'))


class LimitsTests(TestCase):
    """CO-216: no limit on what a file could unpack or read to."""

    def test_a_file_that_unpacks_too_far_is_refused(self):
        from docx import Document
        path = a_file('.docx')
        Document().save(path)
        with mock.patch.object(dpa, 'MAX_UNPACKED_BYTES', 10):
            ok, _text, _pages, error = DocumentProcessingAgent()._extract_text(path, 'docx')
        self.assertFalse(ok)
        self.assertIn('unpacks to', error)

    def test_a_long_text_is_cut_and_says_so(self):
        path = a_file('.txt', text='word ' * 400)
        agent = DocumentProcessingAgent()
        with mock.patch.object(dpa, 'MAX_TEXT_CHARS', 500):
            ok, text, _pages, _err = agent._extract_text(path, 'txt')
        self.assertTrue(ok)
        self.assertEqual(len(text), 500)
        self.assertIn('Only the first 500 characters', agent.read_note)


class CleanInsightsTests(TestCase):
    """CO-268: an amount sent as an object broke the document's page."""

    def test_objects_become_text(self):
        self.assertEqual(dpa.clean_entities({'amounts': [{'value': 18500000, 'currency': 'PKR'}, 'PKR 650,000'],
                                             'names': 'Ali', 'junk': None}),
                         {'amounts': ['18500000 PKR', 'PKR 650,000'], 'names': ['Ali']})
        self.assertEqual(dpa.clean_items([{'text': 'Fees rise 5%'}, '', None, 'Fees rise 5%']), ['Fees rise 5%'])


class LongSummaryTests(TestCase):
    """CO-212: a long file's summary came from its first few sections only."""

    def test_the_sections_are_condensed_until_they_fit(self):
        agent = DocumentSummarizationAgent()
        sections = [f'**Section {i}:** ' + 'x' * 3000 for i in range(12)]
        with mock.patch.object(agent, '_call_llm_for_reasoning', return_value='- condensed points') as ai:
            combined = agent._fit_summaries(sections, 'report.pdf', limit=10000)
        self.assertLessEqual(len(combined), 10000)
        self.assertTrue(ai.called)
        self.assertNotIn('x' * 3000, combined)                     # every section went into a merge, none dropped


class DeadlineTests(OpsCase):
    """CO-277: deadlines were loose sentences, so no deadline alert could fire."""

    def test_deadlines_keep_their_date(self):
        self.assertEqual(clean_deadlines([{'date': '2026-11-15', 'description': 'Renewal notice'}, 'Board review',
                                          {'date': 'null', 'description': 'Audit'}]),
                         [{'date': '2026-11-15', 'description': 'Renewal notice'},
                          {'date': None, 'description': 'Board review'}, {'date': None, 'description': 'Audit'}])

    def test_analytics_lists_coming_deadlines_soonest_first(self):
        today = timezone.localdate()
        OperationsDocumentSummary.objects.create(
            company=self.company, original_filename='lease.pdf', rich_summary='x', processing_status='ready',
            deadlines=[{'date': (today + timedelta(days=9)).isoformat(), 'description': 'Later'},
                       {'date': (today - timedelta(days=3)).isoformat(), 'description': 'Past'},
                       {'date': (today + timedelta(days=2)).isoformat(), 'description': 'Soon'},
                       'Some day'])
        body = self.screen.get('/api/operations/analytics?range=30d').json()
        listed = [d['description'] for d in body['data']['upcoming_deadlines']] if 'data' in body else \
            [d['description'] for d in body['upcoming_deadlines']]
        self.assertEqual(listed, ['Soon', 'Later', 'Some day'])


class ReuploadTests(OpsCase):
    """CO-262: a failed file could not be uploaded again ("already exists"), and Retry said to re-upload."""

    def upload(self, name='lease.txt'):
        from django.core.files.uploadedfile import SimpleUploadedFile
        with mock.patch('core.api_key_service.resolve_for_call'), \
                mock.patch('operations_agent.tasks.process_document_in_background'), \
                self.settings(LOCAL_STORAGE_ROOT=tempfile.mkdtemp()):
            return self.screen.post('/api/operations/documents/upload',
                                    {'file': SimpleUploadedFile(name, b'The rent is PKR 40,000.')})

    def test_a_failed_copy_is_replaced(self):
        self.assertEqual(self.upload().status_code, 202)
        OperationsDocument.objects.update(processing_status='failed', processing_error='Could not read it')
        self.assertEqual(self.upload().status_code, 202)
        self.assertEqual(OperationsDocument.objects.get().processing_status, 'pending')

    def test_a_good_copy_is_still_refused(self):
        self.upload()
        self.assertEqual(self.upload().status_code, 409)

    def test_retry_reads_the_stored_file_when_no_text_was_read(self):
        path = a_file('.txt', text='The rent is PKR 40,000.')
        doc = self.placeholder(processing_status='failed', processing_error='timeout')
        OperationsDocument.objects.filter(pk=doc.pk).update(file=path)
        with mock.patch('core.api_key_service.resolve_for_call'), \
                mock.patch('operations_agent.tasks.process_document_in_background') as started:
            response = self.screen.post(f'/api/operations/documents/{doc.pk}/reprocess')
        self.assertEqual(response.status_code, 202, response.content)
        self.assertEqual(started.call_args.kwargs['file_path'], path)


class ReindexKeepsVectorsTests(OpsCase):
    """CO-266: re-indexing with the embedding provider down replaced good vectors with none."""

    def test_the_old_vectors_stay(self):
        doc = self.placeholder(processing_status='ready', is_processed=True, parsed_text='The rent is PKR 40,000. ' * 60,
                               is_indexed=True, embedding_model='text-embedding-3-small')
        OperationsDocumentChunk.objects.create(document=doc, chunk_index=0, content='The rent', embedding=[0.1, 0.2])
        with mock.patch.object(dpa, 'build_embedding_service', return_value=None):
            result = tasks._reindex_impl(doc.pk)
        self.assertEqual(result['status'], 'not_embedded')
        self.assertEqual(OperationsDocumentChunk.objects.get(document=doc).embedding, [0.1, 0.2])
        doc.refresh_from_db()
        self.assertEqual(doc.processing_status, 'ready')


class ChartDaysTests(OpsCase):
    """CO-267: today's activity was drawn under yesterday's date."""

    def test_the_last_bar_is_today(self):
        self.placeholder(processing_status='ready')
        body = self.screen.get('/api/operations/analytics?range=7d').json()
        series = body.get('data', body)['timeseries']['documents']
        self.assertEqual(len(series), 7)
        self.assertEqual(series[-1]['label'], timezone.localdate().strftime('%b %d'))
        self.assertEqual(series[-1]['value'], 1)


class RealPagesTests(TestCase):
    """CO-278: page numbers were guessed from position."""

    def test_a_passage_gets_the_page_it_starts_on(self):
        page_one, page_two = 'a' * 1500, 'b' * 300
        text = page_one + '\n\n' + page_two
        pairs = DocumentProcessingAgent()._chunk_pairs(text, 2, 'other', [[0, 1], [len(page_one) + 2, 2]])
        self.assertEqual(pairs[0][2], 1)
        self.assertEqual(pairs[-1][2], 2)

    def test_no_pages_no_page_number(self):
        pairs = DocumentProcessingAgent()._chunk_pairs('Some memo text. ' * 100, 1, 'other', None)
        self.assertIsNone(pairs[0][2])


class LiveRunFixesTests(OpsCase):
    """Found on the Operations live run on Margalla, 28 Oct 2026 (F2 to F9)."""

    def test_a_scan_says_why_and_what_to_do(self):                                  # F2
        self.assertIn('scan or a photo', dpa.no_text_message('pdf'))
        self.assertIn('no text in it', dpa.no_text_message('txt'))

    def test_the_summary_asks_for_what_must_be_done_and_amounts_as_written(self):   # F3, F4
        import inspect
        from operations_agent.agents import summarization_agent as sa
        source = inspect.getsource(sa)
        self.assertIn('notices to give, payments due, renewals', source)
        self.assertIn('with its currency', source)

    def test_the_deadline_alert_reads_a_date_and_rings_the_bell(self):               # F5, F9
        from operations_agent.signals import alert_upcoming_deadlines
        from project_manager_agent.models import PMNotification
        today = timezone.localdate()
        summary = OperationsDocumentSummary.objects.create(
            company=self.company, created_by=self.nadia, original_filename='lease.pdf', rich_summary='x',
            processing_status='processing',
            deadlines=[{'date': (today + timedelta(days=5)).isoformat(), 'description': 'Renewal notice'}])
        created = alert_upcoming_deadlines(summary, today)
        self.assertIn('by ' + (today + timedelta(days=5)).strftime('%d %b %Y').lstrip('0'), created.message)
        bell = PMNotification.objects.filter(company_user=self.nadia, data__kind='operations_deadline').first()
        self.assertIsNotNone(bell)
        self.assertIn('lease.pdf', bell.title)

    def test_an_english_follow_up_gets_no_language_line(self):                       # F6
        from core.answer_language import line
        self.assertEqual(line('On the results slide, what about the margin?',
                              'What does the speaker say about the margin on the Q3 results slide?', 'english'), '')
        self.assertIn('Roman Urdu', line('margin kitna hai?', 'What is the margin?', 'roman_urdu'))
        self.assertIn('Urdu script', line('دفتر کے اوقات؟', 'What are the office hours?', 'urdu'))

    def test_company_matters_in_any_script_are_documents(self):                      # F7
        self.assertIn('in any language or script', intent._PROMPT)

    def test_a_real_title_and_todays_date(self):                                     # F8
        from operations_agent.agents.document_authoring_agent import DocumentAuthoringAgent
        agent = DocumentAuthoringAgent()
        with mock.patch.object(agent, '_call_llm_for_reasoning', return_value='Q3 Results and FinEdge Renewal'):
            self.assertEqual(agent._suggest_title('A one-page memo to the management team before ...', 'Memo'),
                             'Q3 Results and FinEdge Renewal')
        with mock.patch.object(agent, '_call_llm_for_reasoning', side_effect=ValueError('down')):
            self.assertTrue(agent._suggest_title('A memo about the lease', 'Memo').startswith('A memo about'))
        with mock.patch.object(agent, '_suggest_title', return_value='Memo'):
            prepared = agent._prepare_prompt(self.company.id, 'memo', 'memo', 'formal', None, None)
        self.assertIn("TODAY'S DATE: ", prepared['user_prompt'])


class VersionsTests(OpsCase):
    """CO-213: regenerating or saving overwrote the old text for good."""

    def test_an_earlier_version_can_be_restored(self):
        doc = OperationsGeneratedDocument.objects.create(company=self.company, generated_by=self.nadia, title='Memo',
                                                         prompt='p', content='First text', version=1, edit_history=[])
        self.screen.patch(f'/api/operations/authoring/documents/{doc.pk}/update', json.dumps({'content': 'Second'}),
                          content_type='application/json')
        listed = self.screen.get(f'/api/operations/authoring/documents/{doc.pk}').json()
        entry = listed['document']['edit_history'][0]
        self.assertTrue(entry['restorable'])
        self.assertNotIn('previous_content', entry)                 # old texts stay out of the responses
        restored = self.screen.post(f'/api/operations/authoring/documents/{doc.pk}/restore',
                                    json.dumps({'version': 1}), content_type='application/json').json()
        self.assertEqual(restored['document']['content'], 'First text')
        self.assertEqual(restored['document']['version'], 3)
